"""
frascold_model.py

Pure Frascold compressor-map model.

PURPOSE
-------
This file only loads the Frascold Excel map and returns compressor performance.
It does NOT create plots and it does NOT export result files.

Inputs:
    T_evap_C
    T_cond_C
    frequency_Hz

Outputs:
    Q_evap
    Q_cond
    compressor power
    current
    mass flow
    discharge temperature
    cooling EER
    heating COP

Extrapolation is disabled.
"""

# ============================================================
# 1. IMPORTS
# ============================================================

from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Optional, Tuple

import numpy as np
from openpyxl import load_workbook


# ============================================================
# 2. MODEL CONSTANTS
# ============================================================

EXPECTED_COMPRESSOR = "V30-84AXHT"
EXPECTED_REFRIGERANT = "R290"

FREQUENCY_SHEET_RE = re.compile(
    r"^F_(\d+(?:\.\d+)?)Hz$",
    re.IGNORECASE,
)

OUTPUT_FIELDS = (
    "evaporator_capacity_kW",
    "condenser_capacity_kW",
    "compressor_power_kW",
    "current_A",
    "mass_flow_kg_h",
    "discharge_temperature_C",
)


# ============================================================
# 3. DATA STRUCTURES
# ============================================================

@dataclass(frozen=True)
class GridPoint:
    frequency_Hz: float
    T_evap_C: float
    T_cond_C: float
    evaporator_capacity_kW: float
    condenser_capacity_kW: float
    compressor_power_kW: float
    current_A: float
    mass_flow_kg_h: float
    discharge_temperature_C: float


# ============================================================
# 4. EXCEL PARSER HELPERS
# ============================================================

def _norm(text: object) -> str:
    if text is None:
        return ""

    value = (
        str(text)
        .strip()
        .lower()
        .replace("°", "deg")
    )

    return re.sub(
        r"[^a-z0-9]+",
        "",
        value,
    )


def _to_float(value: object) -> Optional[float]:
    if value is None:
        return None

    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
    ):
        number = float(value)

        return (
            number
            if math.isfinite(number)
            else None
        )

    value = str(value).strip()

    if value in {
        "",
        "-",
        "–",
        "—",
        "nan",
        "None",
    }:
        return None

    value = (
        value
        .replace("\u00a0", "")
        .replace(" ", "")
        .replace(",", ".")
    )

    try:
        number = float(value)

        return (
            number
            if math.isfinite(number)
            else None
        )

    except ValueError:
        return None


def _find_label_value(
    worksheet,
    label: str,
    max_rows: int = 30,
) -> Optional[object]:

    target = _norm(label)

    for row in range(
        1,
        min(
            worksheet.max_row,
            max_rows,
        )
        + 1,
    ):
        for column in range(
            1,
            min(
                worksheet.max_column,
                6,
            )
            + 1,
        ):
            if (
                _norm(
                    worksheet.cell(
                        row,
                        column,
                    ).value
                )
                == target
            ):
                next_column = (
                    column + 1
                )

                if (
                    next_column
                    <= worksheet.max_column
                ):
                    return worksheet.cell(
                        row,
                        next_column,
                    ).value

    return None


def _find_table_header(
    worksheet,
) -> tuple[
    int,
    Dict[str, int],
    Dict[int, str],
]:

    for row in range(
        1,
        min(
            worksheet.max_row,
            80,
        )
        + 1,
    ):
        headers = {
            column:
                (
                    str(
                        worksheet.cell(
                            row,
                            column,
                        ).value
                    ).strip()
                    if worksheet.cell(
                        row,
                        column,
                    ).value is not None
                    else ""
                )
            for column in range(
                1,
                min(
                    worksheet.max_column,
                    20,
                )
                + 1,
            )
        }

        norms = {
            column:
                _norm(value)
            for column, value
            in headers.items()
        }

        has_te = any(
            "tevap" in value
            for value in norms.values()
        )

        has_tc = any(
            "tcond" in value
            for value in norms.values()
        )

        has_power = any(
            "powerinput" in value
            for value in norms.values()
        )

        if not (
            has_te
            and has_tc
            and has_power
        ):
            continue

        mapping: Dict[str, int] = {}

        for column, value in (
            norms.items()
        ):
            if "tevap" in value:
                mapping[
                    "T_evap_C"
                ] = column

            elif "tcond" in value:
                mapping[
                    "T_cond_C"
                ] = column

            elif (
                "evaporatorcapacity"
                in value
            ):
                mapping[
                    "evaporator_capacity_kW"
                ] = column

            elif (
                "refrigeratingcapacity"
                in value
                or "refrigerationcapacity"
                in value
            ):
                mapping[
                    "refrigerating_capacity_kW"
                ] = column

            elif (
                "condensercapacity"
                in value
            ):
                mapping[
                    "condenser_capacity_kW"
                ] = column

            elif (
                "powerinput"
                in value
            ):
                mapping[
                    "compressor_power"
                ] = column

            elif value.startswith(
                "current"
            ):
                mapping[
                    "current_A"
                ] = column

            elif (
                "massflow"
                in value
            ):
                mapping[
                    "mass_flow_kg_h"
                ] = column

            elif (
                "dischargetemperature"
                in value
            ):
                mapping[
                    "discharge_temperature_C"
                ] = column

        required = {
            "T_evap_C",
            "T_cond_C",
            "compressor_power",
        }

        if not required.issubset(
            mapping
        ):
            continue

        if (
            "evaporator_capacity_kW"
            not in mapping
            and
            "refrigerating_capacity_kW"
            not in mapping
        ):
            continue

        return (
            row,
            mapping,
            headers,
        )

    raise ValueError(
        "Could not find the Frascold performance-table header "
        f"in sheet '{worksheet.title}'."
    )


def _sheet_frequency(
    worksheet,
) -> float:

    reported = _to_float(
        _find_label_value(
            worksheet,
            "Frequency:",
        )
    )

    match = FREQUENCY_SHEET_RE.match(
        worksheet.title.strip()
    )

    from_name = (
        float(
            match.group(1)
        )
        if match
        else None
    )

    if (
        reported is not None
        and from_name is not None
        and not math.isclose(
            reported,
            from_name,
            abs_tol=1e-6,
        )
    ):
        raise ValueError(
            f"Sheet '{worksheet.title}' reports "
            f"{reported:g} Hz but the sheet name "
            f"implies {from_name:g} Hz."
        )

    if reported is not None:
        return reported

    if from_name is not None:
        return from_name

    raise ValueError(
        "Could not determine frequency for "
        f"sheet '{worksheet.title}'."
    )


# ============================================================
# 5. FRASCOLD MAP MODEL
# ============================================================

class FrascoldMap:
    def __init__(
        self,
        workbook_path: str | Path,
    ):
        self.workbook_path = (
            Path(workbook_path)
            .expanduser()
            .resolve()
        )

        if not self.workbook_path.exists():
            raise FileNotFoundError(
                "Frascold workbook not found:\n"
                f"{self.workbook_path}"
            )

        if (
            self.workbook_path.suffix.lower()
            not in {
                ".xlsx",
                ".xlsm",
            }
        ):
            raise ValueError(
                "Frascold source must be "
                ".xlsx or .xlsm."
            )

        workbook = load_workbook(
            self.workbook_path,
            data_only=True,
            read_only=False,
        )

        frequency_sheets = [
            worksheet
            for worksheet
            in workbook.worksheets
            if FREQUENCY_SHEET_RE.match(
                worksheet.title.strip()
            )
        ]

        if not frequency_sheets:
            workbook.close()

            raise ValueError(
                "No frequency sheets found. "
                "Expected F_30Hz ... F_70Hz."
            )

        self.points: Dict[
            Tuple[
                float,
                float,
                float,
            ],
            GridPoint,
        ] = {}

        self.sheet_metadata = []

        frequencies = set()
        evaporating_temperatures = set()
        condensing_temperatures = set()

        for worksheet in frequency_sheets:
            frequency = _sheet_frequency(
                worksheet
            )

            (
                header_row,
                columns,
                raw_headers,
            ) = _find_table_header(
                worksheet
            )

            compressor = _find_label_value(
                worksheet,
                "Compressor:",
            )

            refrigerant = _find_label_value(
                worksheet,
                "Refrigerant:",
            )

            superheat = _to_float(
                _find_label_value(
                    worksheet,
                    "Suction gas superheating:",
                )
            )

            subcooling = _to_float(
                _find_label_value(
                    worksheet,
                    "Liquid subcooling:",
                )
            )

            if (
                compressor not in {
                    None,
                    "",
                }
                and EXPECTED_COMPRESSOR.lower()
                not in str(
                    compressor
                ).lower()
            ):
                workbook.close()

                raise ValueError(
                    f"Sheet '{worksheet.title}' "
                    f"contains compressor "
                    f"'{compressor}', expected "
                    f"{EXPECTED_COMPRESSOR}."
                )

            if (
                refrigerant not in {
                    None,
                    "",
                }
                and EXPECTED_REFRIGERANT.lower()
                not in str(
                    refrigerant
                ).lower()
            ):
                workbook.close()

                raise ValueError(
                    f"Sheet '{worksheet.title}' "
                    f"contains refrigerant "
                    f"'{refrigerant}', expected "
                    f"{EXPECTED_REFRIGERANT}."
                )

            power_header = (
                raw_headers[
                    columns[
                        "compressor_power"
                    ]
                ]
                .lower()
            )

            normalized_power_header = (
                power_header
                .replace(" ", "")
            )

            power_is_watts = (
                "[w]"
                in normalized_power_header
                or
                "power input [w]"
                in power_header
            )

            valid_count = 0
            invalid_count = 0

            for row in range(
                header_row + 1,
                worksheet.max_row + 1,
            ):
                T_evap = _to_float(
                    worksheet.cell(
                        row,
                        columns[
                            "T_evap_C"
                        ],
                    ).value
                )

                T_cond = _to_float(
                    worksheet.cell(
                        row,
                        columns[
                            "T_cond_C"
                        ],
                    ).value
                )

                if (
                    T_evap is None
                    or T_cond is None
                ):
                    continue

                frequencies.add(
                    float(
                        frequency
                    )
                )

                evaporating_temperatures.add(
                    float(
                        T_evap
                    )
                )

                condensing_temperatures.add(
                    float(
                        T_cond
                    )
                )

                power_raw = _to_float(
                    worksheet.cell(
                        row,
                        columns[
                            "compressor_power"
                        ],
                    ).value
                )

                Q_evap = None

                if (
                    "evaporator_capacity_kW"
                    in columns
                ):
                    Q_evap = _to_float(
                        worksheet.cell(
                            row,
                            columns[
                                "evaporator_capacity_kW"
                            ],
                        ).value
                    )

                if (
                    Q_evap is None
                    and
                    "refrigerating_capacity_kW"
                    in columns
                ):
                    Q_evap = _to_float(
                        worksheet.cell(
                            row,
                            columns[
                                "refrigerating_capacity_kW"
                            ],
                        ).value
                    )

                Q_cond = None

                if (
                    "condenser_capacity_kW"
                    in columns
                ):
                    Q_cond = _to_float(
                        worksheet.cell(
                            row,
                            columns[
                                "condenser_capacity_kW"
                            ],
                        ).value
                    )

                mass_flow = (
                    _to_float(
                        worksheet.cell(
                            row,
                            columns[
                                "mass_flow_kg_h"
                            ],
                        ).value
                    )
                    if "mass_flow_kg_h"
                    in columns
                    else None
                )

                discharge_temperature = (
                    _to_float(
                        worksheet.cell(
                            row,
                            columns[
                                "discharge_temperature_C"
                            ],
                        ).value
                    )
                    if "discharge_temperature_C"
                    in columns
                    else None
                )

                current = (
                    _to_float(
                        worksheet.cell(
                            row,
                            columns[
                                "current_A"
                            ],
                        ).value
                    )
                    if "current_A"
                    in columns
                    else None
                )

                if (
                    power_raw is None
                    or Q_evap is None
                    or mass_flow is None
                    or discharge_temperature is None
                ):
                    invalid_count += 1
                    continue

                compressor_power_kW = (
                    power_raw / 1000.0
                    if power_is_watts
                    else power_raw
                )

                if Q_cond is None:
                    Q_cond = (
                        Q_evap
                        + compressor_power_kW
                    )

                self.points[
                    (
                        float(
                            frequency
                        ),
                        float(
                            T_evap
                        ),
                        float(
                            T_cond
                        ),
                    )
                ] = GridPoint(
                    frequency_Hz=float(
                        frequency
                    ),
                    T_evap_C=float(
                        T_evap
                    ),
                    T_cond_C=float(
                        T_cond
                    ),
                    evaporator_capacity_kW=float(
                        Q_evap
                    ),
                    condenser_capacity_kW=float(
                        Q_cond
                    ),
                    compressor_power_kW=float(
                        compressor_power_kW
                    ),
                    current_A=(
                        float(
                            current
                        )
                        if current is not None
                        else float("nan")
                    ),
                    mass_flow_kg_h=float(
                        mass_flow
                    ),
                    discharge_temperature_C=float(
                        discharge_temperature
                    ),
                )

                valid_count += 1

            self.sheet_metadata.append({
                "sheet":
                    worksheet.title,

                "frequency_Hz":
                    float(
                        frequency
                    ),

                "compressor":
                    compressor,

                "refrigerant":
                    refrigerant,

                "superheat_K":
                    superheat,

                "subcooling_K":
                    subcooling,

                "valid_points":
                    valid_count,

                "invalid_points":
                    invalid_count,
            })

        workbook.close()

        self.frequencies = np.array(
            sorted(
                frequencies
            ),
            dtype=float,
        )

        self.tevaps = np.array(
            sorted(
                evaporating_temperatures
            ),
            dtype=float,
        )

        self.tconds = np.array(
            sorted(
                condensing_temperatures
            ),
            dtype=float,
        )

        if not self.points:
            raise ValueError(
                "No valid Frascold map "
                "points were found."
            )


    # ========================================================
    # 6. INTERPOLATION HELPERS
    # ========================================================

    @staticmethod
    def _bracket(
        axis: np.ndarray,
        value: float,
        name: str,
    ):
        if (
            value < axis[0]
            or value > axis[-1]
        ):
            raise ValueError(
                f"{name}={value:g} is outside "
                "the Frascold map range "
                f"[{axis[0]:g}, {axis[-1]:g}]. "
                "Extrapolation is disabled."
            )

        exact = np.where(
            np.isclose(
                axis,
                value,
                atol=1e-10,
                rtol=0,
            )
        )[0]

        if len(
            exact
        ):
            exact_value = float(
                axis[
                    exact[0]
                ]
            )

            return (
                exact_value,
                exact_value,
                0.0,
            )

        high_index = int(
            np.searchsorted(
                axis,
                value,
                side="right",
            )
        )

        low_index = (
            high_index - 1
        )

        low = float(
            axis[
                low_index
            ]
        )

        high = float(
            axis[
                high_index
            ]
        )

        weight_high = (
            (value - low)
            / (high - low)
        )

        return (
            low,
            high,
            float(
                weight_high
            ),
        )


    @staticmethod
    def _choices(
        low: float,
        high: float,
        weight_high: float,
    ):
        if low == high:
            return [
                (
                    low,
                    1.0,
                )
            ]

        return [
            (
                low,
                1.0
                - weight_high,
            ),
            (
                high,
                weight_high,
            ),
        ]


    # ========================================================
    # 7. PUBLIC MODEL FUNCTIONS
    # ========================================================

    def interpolate(
        self,
        T_evap_C: float,
        T_cond_C: float,
        frequency_Hz: float,
    ) -> dict:

        (
            f0,
            f1,
            wf,
        ) = self._bracket(
            self.frequencies,
            frequency_Hz,
            "frequency_Hz",
        )

        (
            e0,
            e1,
            we,
        ) = self._bracket(
            self.tevaps,
            T_evap_C,
            "T_evap_C",
        )

        (
            c0,
            c1,
            wc,
        ) = self._bracket(
            self.tconds,
            T_cond_C,
            "T_cond_C",
        )

        weighted_points = []
        missing = []

        for frequency, frequency_weight in self._choices(
            f0,
            f1,
            wf,
        ):
            for T_evap, evap_weight in self._choices(
                e0,
                e1,
                we,
            ):
                for T_cond, cond_weight in self._choices(
                    c0,
                    c1,
                    wc,
                ):
                    key = (
                        frequency,
                        T_evap,
                        T_cond,
                    )

                    point = self.points.get(
                        key
                    )

                    if point is None:
                        missing.append(
                            key
                        )

                    else:
                        weighted_points.append(
                            (
                                frequency_weight
                                * evap_weight
                                * cond_weight,
                                point,
                            )
                        )

        if missing:
            raise ValueError(
                "INVALID_FRASCOLD_CELL: "
                "Requested point touches an invalid "
                "('-') Frascold workbook cell."
            )

        result = {
            "frequency_Hz":
                float(
                    frequency_Hz
                ),

            "T_evap_C":
                float(
                    T_evap_C
                ),

            "T_cond_C":
                float(
                    T_cond_C
                ),
        }

        for field in (
            OUTPUT_FIELDS
        ):
            values = [
                (
                    weight,
                    getattr(
                        point,
                        field,
                    ),
                )
                for weight, point
                in weighted_points
            ]

            if any(
                not math.isfinite(
                    value
                )
                for _, value
                in values
            ):
                result[
                    field
                ] = float(
                    "nan"
                )

            else:
                result[
                    field
                ] = float(
                    sum(
                        weight
                        * value
                        for weight, value
                        in values
                    )
                )

        compressor_power = (
            result[
                "compressor_power_kW"
            ]
        )

        result[
            "cooling_EER"
        ] = (
            result[
                "evaporator_capacity_kW"
            ]
            / compressor_power
        )

        result[
            "heating_COP"
        ] = (
            result[
                "condenser_capacity_kW"
            ]
            / compressor_power
        )

        return result


    def predict(
        self,
        T_evap_C: float,
        T_cond_C: float,
        frequency_Hz: float,
        mode: str,
    ) -> dict:

        result = self.interpolate(
            T_evap_C,
            T_cond_C,
            frequency_Hz,
        )

        mode = (
            mode
            .strip()
            .lower()
        )

        if mode == "heating":
            result[
                "mode"
            ] = "heating"

            result[
                "useful_capacity_kW"
            ] = result[
                "condenser_capacity_kW"
            ]

            result[
                "COP_or_EER"
            ] = result[
                "heating_COP"
            ]

        elif mode == "cooling":
            result[
                "mode"
            ] = "cooling"

            result[
                "useful_capacity_kW"
            ] = result[
                "evaporator_capacity_kW"
            ]

            result[
                "COP_or_EER"
            ] = result[
                "cooling_EER"
            ]

        else:
            raise ValueError(
                "mode must be "
                "'heating' or 'cooling'."
            )

        return result
